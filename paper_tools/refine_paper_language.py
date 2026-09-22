from __future__ import annotations

import copy
import shutil
import tempfile
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
PAPER_DIR = ROOT / "论文初稿"
SOURCE = PAPER_DIR / "毛绒陪伴机器人语音与动作系统设计_正式版_图示优化版.docx"
OUTPUT = PAPER_DIR / "毛绒陪伴机器人语音与动作系统设计_正式版_语言精修版.docx"


PARAGRAPH_REPLACEMENTS = {
    34: "面向日常陪伴场景中用户对自然语音交互、情绪回应和具身动作反馈的需求，本文设计并实现了一套毛绒陪伴机器人语音与动作系统。系统以 ESP32-S3 为本地感知与执行终端，连接 INMP441 数字麦克风、MAX98357A 数字功放、扬声器、PCA9685 舵机驱动板及四个脚部舵机和一个尾巴舵机；电脑端 AI bridge 负责语音识别、对话生成和语音合成。用户说出唤醒词后，设备进行端点检测录音，并通过 Wi-Fi 将 16 kHz PCM 语音上传至服务端；服务端依次完成 faster-whisper 语音识别、基于“糯糯”角色提示词的 Ollama 对话生成及 TTS 合成，返回回复文本、动作意图、耗时和音频地址。ESP32 下载并播放 WAV 音频，同时将动作意图投入 FreeRTOS 动作队列，由 PCA9685 输出五路 PWM 信号。开发过程中，系统完成了录音方式、识别参数、角色配置、语音合成、触摸输入和舵机供电方案的改进。当前语音链路和舵机通道0已经完成实际验证，下一阶段将围绕五舵机协同运行、连续对话和量化评价展开。",
    38: "This study designs and implements a voice and motion system for a plush companion robot. An ESP32-S3 serves as the local sensing and execution terminal and connects to an INMP441 digital microphone, a MAX98357A audio amplifier, a loudspeaker, a PCA9685 servo driver, four leg servos and one tail servo. A computer-side AI bridge performs speech recognition, dialogue generation and speech synthesis. After wake-word detection, the device records speech through endpoint detection and uploads 16 kHz PCM audio through Wi-Fi. The server uses faster-whisper for speech recognition, Ollama with the Nuonuo persona prompt for dialogue generation, and SAPI or GPT-SoVITS for speech synthesis. It returns the reply text, motion intent, timing information and an audio URL. The ESP32 downloads and plays the WAV audio while dispatching motion intents to a FreeRTOS queue. The voice pipeline and servo channel 0 have completed physical verification. The next stage focuses on coordinated five-servo operation, continuous dialogue and quantitative evaluation.",
    49: "本文重点讨论原型系统的软硬件设计、实现过程和测试方法。当前代码已经配置前左、前右、后左、后右和尾巴五个舵机通道，并实现动作队列与同步插值逻辑。实际硬件验证覆盖 PCA9685 通道0的一只 SG90 舵机，因此本文将“五路控制代码”和“通道0硬件验证”分别表述，确保实现状态与实验依据一致。",
    50: "服务端采用 faster-whisper small、Ollama qwen2.5:3b 和 Windows SAPI TTS，并完成 GPT-SoVITS 独立接口联调。连续对话固件已经实现“30秒无语音退出、最多8轮”的控制逻辑并通过编译。下一阶段将完成新版固件烧录、五舵机联调和标准语料评价，使软件实现、硬件表现和量化结果形成对应关系。",
    62: "硬件部分以 ESP32-S3 为核心。INMP441 是数字 MEMS 麦克风，使用 I2S 输出音频数据；MAX98357A 是 I2S 数字功放，将 ESP32 输出的数字音频转换为可驱动喇叭的信号；PCA9685 是16通道 PWM 驱动器，通过 I2C 接收控制命令后输出舵机脉冲。系统使用 PCA9685 的前五个通道分别连接四个脚部舵机和一个尾巴舵机，使各肢体能够独立校准并执行组合动作。",
    68: "PCA9685 通道0至4依次分配给前左、前右、后左、后右和尾巴。脚部角度范围设置为70°至110°，尾巴设置为60°至120°，主页均为90°。这些参数来自空载台架的安全范围，下一步将结合毛绒骨架、舵机安装方向和机械干涉进行标定。软件会拒绝超出范围的串口命令，降低撞限位和机械卡滞的风险。",
    74: "图3-1  成品供电结构",
    77: "设备上电后初始化音频输入输出、唤醒模型、命令词模型、Wi-Fi 和动作模块，并进入等待唤醒状态。检测到“你好小智”后，系统播放本地欢迎音；AI bridge 连接正常时进入端点检测录音，服务不可用时保留本地命令词模式。开发过程中逐步完善了状态日志和前验检查，现在能够显示 Wi-Fi 状态、服务器地址、AI bridge ready 状态及失败原因，从而明确区分网络、服务和录音环节的问题。",
    81: "当前 ASR 使用 faster-whisper small、CPU int8、beam_size=5、best_of=5，并关闭跨片段文本条件；识别结果通过 OpenCC 转换为简体。四段旧录音的复查结果表明，small 模型能够改善部分错句。服务端同时提供 benchmark_asr.py 脚本，下一步计划建立带人工参考文本的标准语料，统计 CER、整句正确率和识别耗时。",
    84: "TTS 服务提供 Windows SAPI 和 GPT-SoVITS 两个后端，输出统一转换为 ESP32 可播放的16 kHz、16 bit、单声道 WAV。当前 AI bridge 使用 Windows SAPI 完成稳定输出，GPT-SoVITS 调试接口已经生成规范 WAV。下一阶段将把自定义音色接入 ESP32 整机链路，比较两种后端的播放效果。ESP32 下载完整 WAV 后解析 data 区块，并通过 I2S 送往 MAX98357A 播放。",
    86: "服务端 motion 字段支持 happy、shy、comfort、curious 和 none。ESP32 在机器控制开关开启后，把相应意图投递到长度为8的 FreeRTOS 队列。动作任务根据目标姿态与当前姿态之间的最大角度差确定插值步数，并同步更新五个通道，避免网络、录音或播放任务被长时间动作阻塞。第一版中 comfort 和 curious 共用 curious 动作；下一版将结合机器人外观细化情绪与动作的对应关系。",
    94: "配置入口与实际运行效果分开管理。GPT-SoVITS 可以通过配置启用，触摸输入也保留 GPIO 和动作映射；当前默认运行组合为 Windows SAPI 和关闭触摸输入。论文据此区分当前运行配置与下一阶段扩展内容，使个性化能力的描述与系统状态保持一致。",
    95: "6 开发过程与改进",
    96: "6.1 录音与网络链路改进",
    97: "项目早期采用固定5秒录音，短句会产生多余静音和等待，长句或停顿又容易被截断。之后改用端点检测，根据环境噪声基线、动态阈值、开口确认、pre-roll 和连续静音自动确定录音边界，同时保留最长录音时长作为保护。该方案已经写入当前代码，下一步计划采集更多短句、长句和停顿样本，形成完整率更高的实验结论。",
    98: "网络链路早期采用静默回退，日志信息较少，难以判断问题位于网络、服务器还是录音模块。开发过程中逐步完善了日志内容和前验保护，在每次唤醒前检查 Wi-Fi 与 AI bridge 状态，并记录服务器地址和失败原因，使故障位置能够被快速识别。",
    99: "6.2 语音识别、回复与合成改进",
    100: "早期 ASR 使用 tiny 模型，速度较快，但复杂语句和含糊录音的识别效果有限。系统随后升级为 small 模型，并调整 beam_size、best_of、跨片段文本条件和繁简转换。完整提示词和强热词容易在短录音中产生幻听，因此当前默认关闭，只保留配置入口。旧录音复查显示识别结果有所改善，下一步计划用标准语料形成统一的准确率和耗时评价。",
    101: "回复生成最初以规则话术为主，之后加入 Ollama 本地模型、独立角色提示词、短期上下文、情绪字段和动作字段，使机器人形象能够通过配置切换。语音合成以 Windows SAPI 作为默认后端，同时完成 GPT-SoVITS 参数适配和 WAV 格式统一，为自定义音色接入提供接口。",
    102: "6.3 动作与硬件控制改进",
    103: "早期动作控制通道较少，并采用阻塞式调用，容易影响录音、网络和播放任务。系统随后扩展为五通道配置，并使用独立 FreeRTOS 动作队列完成同步插值。开发中还发现未连接的触摸引脚会因悬空而反复触发动作，因此默认关闭触摸输入，并在关闭机器控制时清空动作队列和复位姿态。",
    104: "修改后，PCA9685 通道0的 servo fl 80 和 servo fl 100 指令均能正常执行，动作只入队一次，说明通道0、动作队列和触摸输入保护能够协同工作。下一阶段将按照通道1至4、双舵机协同和五舵机整机的顺序扩展验证范围。",
    105: "6.4 阶段结果与下一步计划",
    107: "目前已经形成语音链路日志、通道0舵机运动记录、四段旧录音对比结果和 GPT-SoVITS 接口输出。下一步计划建立20～50条带人工标准文本的语料，统计唤醒、识别、端点检测和分阶段响应时间；硬件部分将依次完成其余通道、双舵机和五舵机协同实验，并记录峰值电流、温升、异常复位、动作噪声和机械卡滞情况。",
    110: "在软件实现方面，系统已经具备端点检测录音、PCM 上传、结构化服务端返回、动态 WAV 下载播放、角色提示词加载和动作队列控制等能力。开发过程围绕三条主线展开：录音方式由固定时长改为端点检测，AI 服务链路增加状态日志和前验检查，动作控制由阻塞调用改为五通道队列与同步插值。ASR 模型、角色配置、TTS 格式和触摸输入也完成了相应调整。",
    111: "当前语音链路和 PCA9685 通道0已经完成实际验证，五通道控制、连续对话和自定义音色具备代码与接口基础。下一阶段将完成五舵机与电源系统联调、连续对话固件烧录、标准 ASR 语料评价和 GPT-SoVITS 整机播放，并通过重复实验形成识别准确率、响应延迟、运行稳定性和陪伴体验的量化结论。",
}


TABLE_REPLACEMENTS = {
    (0, 2, 2): "已写入代码；下一步形成量化结果",
    (0, 4, 2): "已完成代码与编译；下一步完成烧录验证",
    (0, 6, 2): "下一阶段开展并发、负载和温升实验",
    (1, 5, 2): "通道0已验证；下一步扩展其余通道",
    (1, 6, 2): "采用2S电池和5V星形分支",
    (3, 2, 3): "下一步验证",
    (3, 3, 3): "下一步验证",
    (3, 4, 3): "下一步验证",
    (3, 5, 3): "下一步验证",
    (4, 6, 2): "已提供入口；下一步结合机械结构校准",
    (5, 0, 2): "阶段结论",
    (5, 1, 2): "已完成通道0单机实验",
    (5, 2, 2): "已完成代码验证；下一步开展多机实验",
    (5, 3, 2): "已形成历史链路证据",
    (5, 4, 2): "为模型选型提供依据；下一步建立标准CER",
    (5, 5, 2): "已完成接口联调；下一步完成整机播放",
    (5, 6, 2): "已完成代码与编译；下一步完成烧录验证",
}


def replace_paragraph(paragraph, text: str) -> None:
    first_rpr = None
    if paragraph.runs and paragraph.runs[0]._element.rPr is not None:
        first_rpr = copy.deepcopy(paragraph.runs[0]._element.rPr)
    p = paragraph._element
    for child in list(p):
        if child.tag.endswith("}pPr"):
            continue
        p.remove(child)
    run = paragraph.add_run(text)
    if first_rpr is not None:
        run._element.insert(0, first_rpr)


def patch_images(docx_path: Path) -> None:
    replacements = {
        "word/media/image3.png": PAPER_DIR / "图2-1_四层总体架构.png",
        "word/media/image4.png": PAPER_DIR / "图3-1_成品供电结构.png",
        "word/media/image5.png": PAPER_DIR / "图4-1_语音交互数据流.png",
    }
    with tempfile.TemporaryDirectory(prefix="paper_language_") as temp_dir:
        patched = Path(temp_dir) / docx_path.name
        with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(patched, "w") as zout:
            for item in zin.infolist():
                data = replacements[item.filename].read_bytes() if item.filename in replacements else zin.read(item.filename)
                zout.writestr(item, data)
        shutil.copy2(patched, docx_path)


def main() -> None:
    document = Document(SOURCE)
    for index, text in PARAGRAPH_REPLACEMENTS.items():
        replace_paragraph(document.paragraphs[index], text)
    for (table_index, row_index, cell_index), text in TABLE_REPLACEMENTS.items():
        replace_paragraph(document.tables[table_index].cell(row_index, cell_index).paragraphs[0], text)
    document.save(OUTPUT)
    patch_images(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
