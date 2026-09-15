from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "论文初稿"
OUT_DOCX = OUT_DIR / "毛绒陪伴机器人语音与动作系统设计_初版.docx"


def font(size, bold=False):
    paths = [r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\simhei.ttf"]
    for path in paths:
        if Path(path).is_file():
            return ImageFont.truetype(path, size=size, index=0)
    return ImageFont.load_default()


def arrow(draw, start, end, fill=(70, 70, 70), width=4):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = end
    x0, y0 = start
    if abs(x1 - x0) >= abs(y1 - y0):
        sign = 1 if x1 > x0 else -1
        points = [(x1, y1), (x1 - 16 * sign, y1 - 9), (x1 - 16 * sign, y1 + 9)]
    else:
        sign = 1 if y1 > y0 else -1
        points = [(x1, y1), (x1 - 9, y1 - 16 * sign), (x1 + 9, y1 - 16 * sign)]
    draw.polygon(points, fill=fill)


def box(draw, xy, title, detail, fill, title_font, body_font):
    x0, y0, x1, y1 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=(85, 85, 85), width=2)
    draw.multiline_text((x0 + 16, y0 + 15), title, font=title_font, fill=(20, 20, 20), spacing=6)
    draw.multiline_text((x0 + 16, y0 + 55), detail, font=body_font, fill=(45, 45, 45), spacing=5)


def make_architecture(path):
    image = Image.new("RGB", (1700, 960), "white")
    draw = ImageDraw.Draw(image)
    title_font, body_font = font(33, True), font(23)
    draw.text((40, 24), "图 2-1  毛绒陪伴机器人四层总体架构（初版）", font=title_font, fill=(0, 0, 0))
    columns = [
        (60, "用户交互层", "用户说话、唤醒\n听取回复、观察动作", (254, 236, 218)),
        (455, "设备感知与执行层", "ESP32-S3、INMP441\nMAX98357A、喇叭\nPCA9685、五路舵机", (222, 235, 247)),
        (885, "通信与业务调度层", "Wi-Fi、HTTP、FastAPI\nPCM/WAV、JSON\n调度与结果汇总", (228, 242, 226)),
        (1295, "AI能力层", "faster-whisper、Ollama\n糯糯提示词、TTS\n情绪与动作意图", (237, 226, 244)),
    ]
    for x, label, content, color in columns:
        draw.rounded_rectangle((x, 150, x + 330, 680), radius=20, fill=color, outline=(90, 90, 90), width=2)
        draw.text((x + 20, 180), label, font=font(27, True), fill=(0, 0, 0))
        draw.multiline_text((x + 24, 280), content, font=font(24), fill=(25, 25, 25), spacing=10)
    for x in (390, 820, 1230):
        arrow(draw, (x, 415), (x + 55, 415))
    draw.text((410, 355), "声音/动作", font=font(19), fill=(65, 65, 65))
    draw.text((840, 355), "16 kHz PCM / JSON", font=font(19), fill=(65, 65, 65))
    draw.text((1240, 355), "识别文字 / 回复", font=font(19), fill=(65, 65, 65))
    draw.rounded_rectangle((120, 760, 1580, 885), radius=18, fill=(246, 246, 246), outline=(110, 110, 110), width=2)
    draw.text((155, 788), "基础设施：2S 电池 → 保护 / 开关 → 5V 稳压 → 星形分支；PCA9685 VCC 为逻辑电源，V+ 为舵机大电流电源；所有模块共地。", font=font(23), fill=(20, 20, 20))
    image.save(path)


def make_dataflow(path):
    image = Image.new("RGB", (1700, 1120), "white")
    draw = ImageDraw.Draw(image)
    title_font, body_font = font(33, True), font(21)
    draw.text((40, 24), "图 4-1  语音交互与动作执行数据流（初版）", font=title_font, fill=(0, 0, 0))
    boxes = [
        ((70, 150, 390, 285), "本地唤醒", "WakeNet9 检测\n“你好小智”", (254, 236, 218)),
        ((470, 150, 790, 285), "端点检测录音", "噪声基线、pre-roll\n16 kHz PCM", (222, 235, 247)),
        ((870, 150, 1190, 285), "AI bridge", "HTTP 上传、PCM→WAV\nASR 与调度", (228, 242, 226)),
        ((1270, 150, 1590, 285), "对话与合成", "Ollama→reply/motion\nTTS→audio_url", (237, 226, 244)),
        ((870, 510, 1190, 655), "ESP32 播放", "下载 WAV、解析 data\nMAX98357A→喇叭", (222, 235, 247)),
        ((1270, 510, 1590, 655), "动作队列", "motion→FreeRTOS\nPCA9685→五路 PWM", (254, 236, 218)),
    ]
    for xy, head, detail, color in boxes:
        box(draw, xy, head, detail, color, title_font, body_font)
    arrow(draw, (390, 218), (470, 218)); draw.text((395, 185), "欢迎音后", font=font(18), fill=(70, 70, 70))
    arrow(draw, (790, 218), (870, 218)); draw.text((800, 185), "HTTP：PCM", font=font(18), fill=(70, 70, 70))
    arrow(draw, (1190, 218), (1270, 218)); draw.text((1195, 185), "识别文字", font=font(18), fill=(70, 70, 70))
    arrow(draw, (1430, 285), (1430, 510)); draw.text((1450, 370), "JSON：reply_text\nmotion/audio_url", font=font(18), fill=(70, 70, 70))
    arrow(draw, (1350, 510), (1170, 655)); draw.text((1165, 605), "audio_url", font=font(18), fill=(70, 70, 70))
    draw.rounded_rectangle((120, 830, 1580, 955), radius=18, fill=(248, 248, 248), outline=(115, 115, 115), width=2)
    draw.multiline_text((150, 855), "状态说明：端点检测、服务端管线与 WAV 下载播放均已有代码；历史日志存在端到端请求记录。\n“最多 8 轮、30 秒无语音退出”的最新连续对话配置已编译，尚未烧录实测。", font=font(22), fill=(25, 25, 25), spacing=8)
    image.save(path)


def make_power(path):
    image = Image.new("RGB", (1700, 880), "white")
    draw = ImageDraw.Draw(image)
    draw.text((40, 24), "图 3-1  成品供电结构（设计方案，待整机验证）", font=font(33, True), fill=(0, 0, 0))
    box(draw, (80, 280, 350, 430), "2S 7.4V 电池", "满电约 8.4V\n需匹配保护", (254, 236, 218), font(28, True), font(21))
    box(draw, (450, 280, 720, 430), "保护 / 开关", "BMS、保险丝或\n自恢复保险丝", (237, 226, 244), font(28, True), font(21))
    box(draw, (820, 280, 1090, 430), "5V 稳压模块", "建议 5V/6A～8A\n按实测峰值校核", (228, 242, 226), font(28, True), font(21))
    box(draw, (1230, 110, 1590, 245), "舵机支路", "PCA9685 V+\n五只舵机；大电容", (255, 232, 220), font(28, True), font(21))
    box(draw, (1230, 330, 1590, 465), "主控逻辑支路", "ESP32 5V/VIN\n板载3.3V→INMP441\nPCA9685 VCC", (222, 235, 247), font(28, True), font(21))
    box(draw, (1230, 585, 1590, 720), "音频支路", "MAX98357A VIN\n按模块额定电压连接", (222, 235, 247), font(28, True), font(21))
    arrow(draw, (350, 355), (450, 355)); arrow(draw, (720, 355), (820, 355)); arrow(draw, (1090, 355), (1230, 175)); arrow(draw, (1090, 355), (1230, 397)); arrow(draw, (1090, 355), (1230, 652))
    draw.text((85, 615), "所有支路 GND 在电源分配点共地；舵机电流不得经过 ESP32 开发板或普通杜邦线。\nUSB 仅用于烧录/调试；外部 5V 与 USB 同时连接前需确认隔离，避免倒灌。", font=font(25), fill=(20, 20, 20))
    image.save(path)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_run_font(run, name="宋体", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_text(paragraph, text, name="宋体", size=10.5, bold=False):
    run = paragraph.add_run(text)
    set_run_font(run, name, size, bold)
    return run


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    add_text(p, text)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_before = Pt(6 if level == 1 else 0)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 0)
    add_text(p, text, "黑体" if level == 1 else "宋体", 12 if level == 1 else 10.5, True)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    add_text(p, text, "宋体", 10.5)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, text in enumerate(headers):
        set_cell_shading(header_cells[index], "D9EAF7")
        header_cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = header_cells[index].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, text, "宋体", 9, True)
    for row in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row):
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[index].paragraphs[0]
            add_text(p, str(text), "宋体", 9)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()
    return table


def add_figure(doc, path, caption, width=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width))
    add_caption(doc, caption)


def add_cover(doc):
    section = doc.sections[0]
    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "项目编号：____________", "宋体", 12)
    for _ in range(4): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "本科生研究计划（PRP）研究论文", "黑体", 22, True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "（第    期）", "宋体", 16)
    for _ in range(7): doc.add_paragraph()
    for label in ["论文题目：毛绒陪伴机器人语音与动作系统设计", "项目负责人：__________    学院（系）：__________", "指导教师：__________    学院（系）：__________", "参与学生：________________________________", "项目执行时间：____年__月 至 ____年__月"]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, label, "楷体", 15)
        p.paragraph_format.space_after = Pt(9)
    doc.add_page_break()


def add_header_footer(section):
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(header, "上海交通大学第    期PRP学生研究论文", "楷体", 9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def add_abstracts(doc):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "摘要", "黑体", 14, True)
    abstract = (
        "面向日常陪伴场景中用户对自然语音交互、情绪回应和具身动作反馈的需求，本文设计并实现了一套毛绒陪伴机器人语音与动作系统。"
        "系统以 ESP32-S3 为本地感知与执行终端，连接 INMP441 数字麦克风、MAX98357A 数字功放、扬声器、PCA9685 舵机驱动板及四个脚部舵机和一个尾巴舵机；"
        "电脑端 AI bridge 负责语音识别、对话生成和语音合成等计算量较大的任务。用户说出唤醒词后，设备进行端点检测录音，并通过 Wi-Fi 将 16 kHz PCM 语音上传至服务端；"
        "服务端依次完成 faster-whisper 语音识别、基于“糯糯”角色提示词的 Ollama 对话生成及 TTS 合成，返回回复文本、动作意图、耗时和音频地址。"
        "ESP32 下载并播放 WAV 音频，同时将动作意图投入 FreeRTOS 动作队列，由 PCA9685 输出五路 PWM 信号。本文还记录了录音端点检测、ASR 参数、角色提示词、TTS 回退、触摸悬空误触发和舵机供电等调试过程。"
        "当前已完成语音链路与单舵机通道0的实际验证；五舵机并发、连续对话新版固件和量化实验仍待进一步测试。"
    )
    add_body(doc, abstract)
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Pt(0)
    add_text(p, "关键词：", "黑体", 12, True); add_text(p, "毛绒机器人，语音交互，ESP32-S3，动作控制，陪伴机器人", "宋体", 10.5)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "ABSTRACT", "Times New Roman", 14, True)
    english = (
        "This study designs a voice and motion system for a plush companion robot. The ESP32-S3 is used as the local sensing and execution terminal, connected to an INMP441 digital microphone, a MAX98357A audio amplifier, a loudspeaker, a PCA9685 servo driver, four leg servos and one tail servo. "
        "A computer-side AI bridge performs speech recognition, dialogue generation and speech synthesis. After a wake word is detected, the device records speech through endpoint detection and uploads 16 kHz PCM audio to the server through Wi-Fi. "
        "The server performs faster-whisper speech recognition, Ollama-based dialogue generation with the Nuonuo persona prompt, and TTS synthesis. It then returns reply text, motion intent, timing information and an audio URL. "
        "The ESP32 downloads and plays the WAV audio while dispatching motion intents to a FreeRTOS queue. At present, the voice pipeline and servo channel 0 have been physically tested. Five-servo concurrent control, the latest continuous-dialogue firmware and quantitative evaluation remain to be completed."
    )
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Pt(21); p.paragraph_format.line_spacing = 1
    add_text(p, english, "Times New Roman", 10.5)
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Pt(0)
    add_text(p, "KEY WORDS：", "Times New Roman", 12, True); add_text(p, "plush robot, voice interaction, ESP32-S3, motion control, companion robot", "Times New Roman", 10.5)
    doc.add_page_break()


def build_document():
    OUT_DIR.mkdir(exist_ok=True)
    arch = OUT_DIR / "图2-1_四层总体架构.png"
    flow = OUT_DIR / "图4-1_语音交互数据流.png"
    power = OUT_DIR / "图3-1_成品供电结构.png"
    make_architecture(arch); make_dataflow(flow); make_power(power)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(3.2); section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.5); section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.5); section.footer_distance = Cm(1.75)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体"); normal.font.size = Pt(10.5)
    add_header_footer(section)
    add_cover(doc)
    add_abstracts(doc)

    add_heading(doc, "1 绪论")
    add_heading(doc, "1.1 研究背景", 2)
    add_body(doc, "毛绒玩具具有柔软、低压力和容易亲近的外观特点，适合承载轻量化的陪伴交互。与只能播放固定音频的玩具相比，带有语音理解、回复语音和动作反馈能力的毛绒机器人，能够在用户表达疲惫、开心、焦虑或希望被陪伴时，给出相对自然的回应，并通过身体动作增强交互感受。")
    add_body(doc, "本项目以“糯糯”毛绒小龙为机器人形象，目标不是构建医疗诊断或心理治疗设备，而是实现一个可进行日常短对话、提供温和陪伴和做出简单肢体反馈的原型系统。系统设计需要同时处理嵌入式端本地唤醒与录音、无线音频传输、电脑端 AI 服务调用、语音播放、五路舵机控制以及移动供电条件下的稳定性。")
    add_heading(doc, "1.2 研究目标与内容", 2)
    add_body(doc, "本研究拟完成一个由嵌入式终端和电脑端 AI bridge 共同组成的语音交互系统。ESP32-S3 负责本地唤醒词检测、麦克风采集、网络通信、回复音频播放和动作执行；电脑端负责语音识别、大语言模型回复和语音合成。这样的分工避免将较大的 ASR、LLM 和 TTS 模型部署到资源有限的单片机中，同时保留机器人本地的声音输入、声音输出和动作能力。")
    add_body(doc, "本文围绕硬件结构、语音数据流、动作调度、个性化配置和测试方法展开。研究重点包括：设计五路舵机的供电和控制方式；实现唤醒、端点检测录音、PCM 上传、ASR、LLM、TTS、WAV 下载播放和动作执行链路；将角色提示词与业务代码解耦；建立对识别准确率、响应耗时和动作稳定性的测试方案。")
    add_heading(doc, "1.3 研究范围与证据边界", 2)
    add_body(doc, "本文重点讨论原型系统的软硬件设计、实现过程和测试方法。当前代码中已经配置了前左、前右、后左、后右和尾巴五个舵机通道，并完成了动作队列与同步插值逻辑；但实际硬件验证目前仅覆盖 PCA9685 通道0的一只 SG90 舵机。因此，文中将严格区分“五路控制代码已实现”和“五舵机整机已稳定运行”。")
    add_body(doc, "当前服务端已配置 faster-whisper small 模型、Ollama qwen2.5:3b 模型和 Windows SAPI TTS；GPT-SoVITS 已完成单独接口联调，但尚未形成其在 ESP32 整机上的稳定播放证据。连续对话的最新“30 秒无语音退出、最多8轮”固件已通过编译，尚未烧录实测。后续章节中的状态描述均以“设计计划、已写入代码、已编译、已烧录、已实测、已有量化数据”作为判断依据。")

    add_heading(doc, "2 系统总体方案")
    add_heading(doc, "2.1 四层架构", 2)
    add_body(doc, "系统按职责分为四层。用户交互层负责用户说话、唤醒机器人、收听回复和观察动作；设备感知与执行层负责音频采集、播放和舵机输出；通信与业务调度层负责 Wi-Fi、HTTP、PCM/WAV 处理和 JSON 结果汇总；AI 能力层负责语音识别、对话生成、角色约束、语音合成和动作意图生成。供电系统不属于其中单独一层，而是支撑设备层全部模块的基础设施。")
    add_figure(doc, arch, "图2-1  毛绒陪伴机器人四层总体架构")
    add_heading(doc, "2.2 终端与服务端分工", 2)
    add_body(doc, "ESP32-S3 侧的任务强调低时延和与物理设备直接相连，包括 WakeNet 本地唤醒、INMP441 音频读取、端点检测、HTTP 上传、WAV 音频播放以及 PCA9685 动作控制。电脑端服务采用 FastAPI 组织 AI bridge，接收原始 PCM 后转换为 WAV，依次调用 ASR、对话服务和 TTS 服务，并以 JSON 返回识别文字、回复、动作、音频地址和耗时。该方式使服务端模型可以独立替换，ESP32 不需要因更换 ASR、LLM 或 TTS 后端而重新设计协议。")
    add_heading(doc, "2.3 已实现状态", 2)
    add_table(doc, ["子系统", "当前实现", "证据状态"], [
        ["本地唤醒", "WakeNet9 与“你好小智”模型", "已写入代码；历史使用"],
        ["录音", "端点检测、噪声基线、pre-roll、静音结束", "已写入代码；待系统量化"],
        ["AI bridge", "PCM 上传、ASR、Ollama、TTS、JSON 返回", "已写入代码；历史端到端日志"],
        ["连续对话", "最多8轮、30秒无语音退出", "已编译，未烧录实测"],
        ["动作控制", "五路配置、队列、同步插值", "已编译；通道0单机实测"],
        ["五舵机整机", "并发、负载和温升", "待连接与测试"],
    ], [3.2, 7.0, 5.0])

    add_heading(doc, "3 硬件系统设计")
    add_heading(doc, "3.1 硬件组成与职责", 2)
    add_body(doc, "硬件部分以 ESP32-S3 为核心。INMP441 是数字 MEMS 麦克风，使用 I2S 输出音频数据；MAX98357A 是 I2S 数字功放，将 ESP32 输出的数字音频转换为可驱动喇叭的信号；PCA9685 是 16 通道 PWM 驱动器，通过 I2C 接收控制命令后输出舵机脉冲。原型方案使用 PCA9685 的前五个通道分别连接四个脚部舵机和一个尾巴舵机，便于后续扩展和独立校准。")
    add_table(doc, ["模块", "作用", "当前状态"], [
        ["ESP32-S3", "本地语音、Wi-Fi、播放和动作调度", "当前统一工程主控"],
        ["INMP441", "采集用户语音并输出 I2S 数据", "GPIO 已由代码固定"],
        ["MAX98357A+喇叭", "播放欢迎音和服务端回复音频", "代码支持 WAV 下载与播放"],
        ["PCA9685", "输出五路舵机 PWM", "通道0已实际验证"],
        ["SG90 或后续选型舵机", "脚部和尾巴姿态动作", "其余四路待验证"],
        ["2S电池、保护、降压", "移动供电与电压转换", "成品方案待整机验证"],
    ], [3.6, 7.4, 4.2])
    add_heading(doc, "3.2 GPIO 与通信连接", 2)
    add_body(doc, "当前固件将 INMP441 的 WS、SCK 和 SD 分别固定为 GPIO4、GPIO5 和 GPIO6；MAX98357A 的 BCLK、LRC 和 DIN 分别使用 GPIO15、GPIO16 和 GPIO7。PCA9685 通过 I2C 与 ESP32-S3 通信，GPIO1 为 SDA，GPIO2 为 SCL，I2C 时钟设为100 kHz。该分配以当前代码为依据，不能沿用早期单舵机直连 GPIO18 的示例。GPIO18 对应的旧 ServoController 文件仍保留在工程中，但不属于当前 PCA9685 五舵机方案。")
    add_table(doc, ["ESP32-S3 引脚", "连接模块引脚", "通信/用途"], [
        ["GPIO4", "INMP441 WS/LRCLK", "I2S 输入字选择"],
        ["GPIO5", "INMP441 SCK/BCLK", "I2S 输入位时钟"],
        ["GPIO6", "INMP441 SD", "I2S 输入数据"],
        ["GPIO7", "MAX98357A DIN", "I2S 输出数据"],
        ["GPIO15", "MAX98357A BCLK", "I2S 输出位时钟"],
        ["GPIO16", "MAX98357A LRC", "I2S 输出左右声道时钟"],
        ["GPIO1", "PCA9685 SDA", "I2C 数据"],
        ["GPIO2", "PCA9685 SCL", "I2C 时钟"],
    ], [3.2, 5.4, 6.6])
    add_heading(doc, "3.3 五路舵机与安全限位", 2)
    add_body(doc, "PCA9685 通道0至4依次分配给前左、前右、后左、后右和尾巴。当前脚部保守角度范围为70°至110°，尾巴为60°至120°，主页均为90°。这些数值来自空载台架安全限制，尚未根据实际毛绒骨架、舵机安装方向和机械干涉完成标定。软件对于超出范围的串口命令直接拒绝，以减少初期测试中撞限位和卡滞的风险。")
    add_table(doc, ["PCA9685通道", "肢体", "当前范围", "状态"], [
        ["0", "前左", "70°～110°", "已单机实测"],
        ["1", "前右", "70°～110°", "待测"],
        ["2", "后左", "70°～110°", "待测"],
        ["3", "后右", "70°～110°", "待测"],
        ["4", "尾巴", "60°～120°", "待测"],
    ], [3.0, 3.4, 4.2, 4.6])
    add_heading(doc, "3.4 供电与共地设计", 2)
    add_body(doc, "PCA9685 板上的 VCC 和 V+ 必须区分。VCC 用于驱动芯片的逻辑供电，应接 ESP32 的3.3V；V+ 是舵机供电排，应连接外部稳压5V。舵机的信号电压需要以 ESP32 为参考，因此 ESP32、PCA9685 和外部5V电源必须共地。共地不是让舵机从 ESP32 取电，而是建立统一的信号参考电位。")
    add_body(doc, "五个舵机启动、换向或受阻时可能出现较大的瞬时电流。若让它们通过 ESP32 开发板的5V或3.3V引脚供电，可能造成板载稳压器过载、压降或复位，也会把电机噪声带入音频链路。因此，成品设计采用2S电池经过保护、保险丝和总开关后进入5V大电流稳压模块，再以星形分支分别供给 PCA9685 V+、ESP32的5V/VIN和音频模块。舵机支路入口应配置大容量低 ESR 电容，电源线需采用能承载电流的线材；普通杜邦线和面包板不适合承载整机舵机电流。")
    add_figure(doc, power, "图3-1  成品供电结构（设计方案，待整机验证）")

    add_heading(doc, "4 软件系统与交互流程")
    add_heading(doc, "4.1 ESP32 语音状态机", 2)
    add_body(doc, "设备上电后初始化音频输入输出、唤醒模型、命令词模型、Wi-Fi 和动作模块，并进入等待唤醒状态。检测到“你好小智”后，系统播放本地欢迎音。若 AI bridge 已连接，则进入端点检测录音；若网络或服务不可用，则保留本地命令词模式作为回退。该回退机制最初只表现为静默切换，容易被误认为端点检测故障，后续已增加 Wi-Fi 状态、服务器地址和 AI bridge ready 状态日志。")
    add_body(doc, "端点检测先采集环境噪声估计 noise_floor，然后在等待开口阶段维护短时 pre-roll 环形缓冲。当 RMS 连续超过动态开始阈值达到确认时间后，系统把 pre-roll 写入正式录音并进入录音状态；在达到最短录音时长后，连续静音达到设定时间则结束；达到最长时长时也会强制结束。录音结果是16 kHz、16 bit、单声道 PCM。")
    add_heading(doc, "4.2 AI bridge 调度", 2)
    add_body(doc, "ESP32 使用 HTTP POST 上传 PCM，并在请求头中传递采样率和音频格式。FastAPI 服务保存原始 PCM，同时进行音量归一化后写入 WAV。服务端顺序调用 faster-whisper、对话服务和 TTS 服务，最后返回 recognized_text、reply_text、reply_emotion、motion、audio_url、timings_ms 等字段。ASR、LLM 和 TTS 存在数据依赖，因此当前链路以顺序调用为主。")
    add_body(doc, "当前 ASR 使用 faster-whisper small、CPU int8、beam_size=5、best_of=5，并关闭跨片段文本条件；识别结果使用 OpenCC 转换为简体。初步的四段旧录音复查显示 small 模型在部分错句上优于原 tiny 模型，但旧样本没有完整人工标准文本，不能据此报告正式字错误率。服务端提供 benchmark_asr.py 脚本，后续可对带参考文本的20～50条语料计算 CER、整句正确率和耗时。")
    add_heading(doc, "4.3 回复生成、TTS 与动作意图", 2)
    add_body(doc, "对话服务默认调用 Ollama 中的 qwen2.5:3b 模型。角色设定存放在独立的 nuonuo_v1.txt 文件中，业务代码通过配置读取，使机器人名称、表达方式和安全边界能够在不修改对话逻辑的情况下切换。模型被要求返回 reply_text、style、emotion 和 motion 字段，服务端再根据用户输入校正情绪与动作，避免负面情绪被误配为 happy。若 Ollama 不可用或返回解析失败，服务端使用规则回复兜底并标记状态。")
    add_body(doc, "TTS 服务提供 Windows SAPI 和 GPT-SoVITS 两个后端。所有输出在服务端统一为 ESP32 可播放的16 kHz、16 bit、单声道 WAV。当前后台 AI bridge 默认采用 Windows SAPI；GPT-SoVITS 已在调试接口生成过规范 WAV，但未完成当前整机的持续联调。ESP32 下载完整 WAV 后解析 data 区块并通过 I2S 送往 MAX98357A 播放。")
    add_heading(doc, "4.4 动作队列与同步插值", 2)
    add_body(doc, "服务端 motion 字段支持 happy、shy、comfort、curious 和 none。ESP32 在机器控制开关开启时把相应意图投递到长度为8的 FreeRTOS 队列。动作任务从队列中取出命令，按照目标姿态和当前姿态之间的最大角度差确定插值步数，再同步更新五个通道，避免网络、录音或播放任务直接执行长时间阻塞动作。当前 comfort 和 curious 在执行端共用 curious 动作，属于第一版动作映射，后续应结合整机外观和用户体验继续细化。")
    add_figure(doc, flow, "图4-1  语音交互与动作执行数据流")

    add_heading(doc, "5 个性化设计")
    add_heading(doc, "5.1 角色与回复风格", 2)
    add_body(doc, "当前“糯糯”人设描述其为温暖、活泼、略带傲娇的毛绒小龙。提示词同时约束回复应使用简体中文、尽量短、适合语音播放，不进行医学诊断或承诺治疗效果；当出现明显危险、自伤或急症内容时，应鼓励用户寻求现实中的可信任人员或专业帮助。服务端当前提供 cute 和 encourage 两种情境风格，并保留最近四轮对话作为短期上下文。")
    add_heading(doc, "5.2 可配置项目", 2)
    add_table(doc, ["可配置内容", "配置位置", "当前状态"], [
        ["角色名称与提示词", "server/prompts 与环境变量", "已提供入口"],
        ["LLM模型、超时与回复风格", "server_config.py / 环境变量", "已提供入口"],
        ["ASR模型、beam、best_of、热词", "server_config.py / 环境变量", "已提供入口；热词默认关闭"],
        ["TTS后端、参考音频与语言", "server_config.py / 环境变量", "已提供入口"],
        ["端点检测阈值与会话时间", "main/network_config.h", "已提供入口"],
        ["舵机通道、角度范围与主页", "main/robot_config.h", "已提供入口，待机械校准"],
        ["触摸输入", "main/robot_config.h", "代码预留，默认关闭"],
    ], [4.0, 6.0, 5.2])
    add_body(doc, "需要注意的是，配置入口不等于已经完成效果验证。例如，GPT-SoVITS 后端可由配置启用，但当前后台服务尚未运行该 API；触摸输入已有 GPIO 和动作映射，但未接硬件时默认关闭。论文中应将这些内容写为“可配置能力”或“待验证扩展”，而不能表述为成品功能。")

    add_heading(doc, "6 开发过程、测试与讨论")
    add_heading(doc, "6.1 录音和网络链路的迭代", 2)
    add_body(doc, "项目早期采用固定5秒录音。该方式在短句场景中会引入额外静音和等待，而长句或停顿场景又可能出现截断。为此，固件保留固定时长函数作为回退，并新增端点检测函数，通过环境噪声基线、动态开始/停止阈值、80 ms 开口确认、300 ms pre-roll、最短录音时长和连续静音结束判定控制上传边界。该功能已经写入当前代码；需要在短句、长句、停顿和不开口场景下进一步采集串口日志和录音样本，才能形成端点完整率结论。")
    add_body(doc, "一次测试中，欢迎音播放后系统直接进入本地命令词模式，表面现象像是端点检测没有执行。排查后发现根因是 ai_bridge_ready 为 false 时触发了静默回退，且旧日志在 ready 判断前就打印了“准备录音”。后续修改为每次唤醒前检查 Wi-Fi 和 AI bridge 状态，并输出服务器地址和失败原因。该问题说明联网系统需要可观测性，不能仅依赖静默回退。")
    add_heading(doc, "6.2 ASR、回复与TTS的迭代", 2)
    add_body(doc, "ASR 从 tiny 调整为 small 后，解码参数修改为 beam_size=5、best_of=5，关闭跨片段文本条件，并增加繁体转简体。完整提示词和过强热词曾在短而含糊的录音中引发幻听，因此当前默认关闭，仅保留受控实验配置入口。非正式旧录音复查显示 small 模型对部分结果有改善，热启动单条识别约1.6～1.8秒；这只是开发观察，不能替代标准语料的统计实验。")
    add_body(doc, "回复生成方面，项目从规则话术逐步增加 Ollama 本地模型、角色提示词、短期上下文、情绪字段和动作字段。人设文件从业务代码拆出后，能够在不重写服务逻辑的情况下切换机器人形象。TTS 方面，Windows SAPI 被保留为稳定回退；GPT-SoVITS 适配了 api_v2 参数和 WAV 规范化，单独调试接口已成功生成可识别的16 kHz WAV，但当前整机端到端效果仍待验证。")
    add_heading(doc, "6.3 动作与硬件问题的迭代", 2)
    add_body(doc, "为适配四脚加尾巴，动作层从早期的少路、阻塞式控制调整为五通道配置和独立 FreeRTOS 动作任务。公开动作函数仅负责入队，执行任务负责插值和 PWM 写入。统一工程的五路版本已经编译，且 v2 固件曾成功烧录；实际复测中，GPIO8、GPIO9、GPIO10 未连接的触摸输入处于悬空状态，持续误触发动作并占满队列，导致正常串口舵机命令看似无效。修复方案是将 TOUCH_INPUTS_ENABLED 默认设为 false，并使关闭机器控制时清空待执行动作和请求复位。")
    add_body(doc, "修复后，用户实际测试了通道0的 servo fl 80 和 servo fl 100 指令，确认动作只入队一次且舵机运动符合预期。这证明了统一工程中 PCA9685 通道0、动作队列和悬空触摸修复的可用性，但不证明其余四通道、组合动作或五机大电流供电已通过。")
    add_heading(doc, "6.4 当前测试结果与待补实验", 2)
    add_table(doc, ["测试项", "当前结果", "论文表述边界"], [
        ["单舵机台架", "通道0正常运动，无尖叫", "已实测；仅单机"],
        ["五路动作代码", "队列与插值编译通过", "不可写成五机已稳定运行"],
        ["历史AI端到端", "存在 PCM 上传、200响应和WAV下载日志", "证明历史链路通，不代表最新固件"],
        ["ASR small", "四段旧录音的非正式改善观察", "不可报告正式CER"],
        ["GPT-SoVITS", "调试接口生成规范WAV", "当前整机联调待完成"],
        ["连续对话v1", "代码与编译完成", "尚未烧录实测"],
    ], [3.5, 6.0, 5.7])
    add_body(doc, "后续实验应首先建立带人工标准文本的20～50条语料，在安静和普通室内噪声、20 cm和50 cm距离下记录唤醒成功率、误唤醒率、CER、整句正确率和端点完整率。响应速度应记录说话结束、上传完成、ASR完成、LLM完成、TTS完成、下载完成、播放开始和动作开始时间，并报告P50与P90。硬件部分应逐通道验证1至4，再从两机并发逐步扩展到五机，记录峰值电流、异常复位、温升、动作噪声和机械卡滞情况。")

    add_heading(doc, "7 结论")
    add_body(doc, "本文完成了毛绒陪伴机器人语音与动作系统的初步设计与实现。系统采用 ESP32-S3 与电脑端 AI bridge 分工：前者负责本地音频采集、唤醒、网络通信、播放和动作执行，后者负责 faster-whisper 识别、Ollama 对话和 TTS 合成。硬件方案明确了 INMP441、MAX98357A、PCA9685 和五路舵机的通信关系，并提出了2S电池经5V稳压后使用星形分支、舵机独立大电流支路和所有模块共地的成品供电结构。")
    add_body(doc, "在软件实现方面，系统已具备端点检测录音、PCM 上传、结构化服务端返回、动态 WAV 下载播放、角色提示词加载和动作队列控制等能力。开发过程中针对固定时长录音、AI bridge 静默回退、ASR 幻听、角色一致性、TTS 格式、触摸悬空误触发和阻塞式动作等问题进行了多轮修改。当前已完成单舵机通道0的实际验证，并保留了历史语音端到端日志。")
    add_body(doc, "本项目仍处于原型迭代阶段。后续需要完成五路舵机与电源系统的实测、连续对话最新版固件的烧录、端点检测的系统化场景测试、标准 ASR 语料评估、GPT-SoVITS 整机联调及用户体验评价。只有在获得重复测试记录后，才能对识别准确率、延迟、稳定性和陪伴体验形成量化结论。")

    add_heading(doc, "参考文献")
    refs = [
        "[1] Espressif Systems. ESP32-S3 Series Datasheet[EB/OL]. https://www.espressif.com/sites/default/files/documentation/esp32-s3_datasheet_en.pdf.",
        "[2] NXP Semiconductors. PCA9685 16-channel, 12-bit PWM Fm+ I2C-bus LED controller[EB/OL]. https://www.nxp.com/docs/en/data-sheet/PCA9685.pdf.",
        "[3] TDK InvenSense. INMP441 High Accuracy Low Power Digital Output MEMS Microphone[EB/OL]. https://invensense.tdk.com/wp-content/uploads/2015/02/INMP441.pdf.",
        "[4] Analog Devices. MAX98357A/MAX98357B PCM Input Class D Audio Power Amplifiers[EB/OL]. https://www.analog.com/media/en/technical-documentation/data-sheets/MAX98357A-MAX98357B.pdf.",
        "[5] Radford A, Kim J W, Xu T, et al. Robust Speech Recognition via Large-Scale Weak Supervision[EB/OL]. arXiv:2212.04356, 2022.",
        "[6] SYSTRAN. faster-whisper: Faster Whisper transcription with CTranslate2[EB/OL]. https://github.com/SYSTRAN/faster-whisper.",
        "[7] Ollama. Ollama Documentation[EB/OL]. https://docs.ollama.com/.",
        "[8] 毛绒治愈机器人PRP项目组. AI Bridge录音端点检测模块说明[Z]. 2026.",
        "[9] 毛绒治愈机器人PRP项目组. 五舵机硬件集成与语音量化评估计划[Z]. 2026.",
    ]
    for ref in refs:
        p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Pt(0); p.paragraph_format.line_spacing = 1
        add_text(p, ref, "宋体", 10.5)

    add_heading(doc, "致谢")
    add_body(doc, "感谢指导教师在课题方向、系统设计和论文撰写方面给予的指导，感谢参与项目测试与讨论的同学提供意见，也感谢在硬件接线、语音录音和功能验证过程中给予帮助的人员。后续将继续完善实验记录，使本研究从可运行原型逐步发展为具有完整证据链的项目成果。")

    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build_document()
